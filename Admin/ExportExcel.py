
import os
from openpyxl import Workbook, load_workbook
from rich import print


def ExportExcel(data, headers, file_name, sheet_name):
    if not data:
        print("[bold yellow]No data available to export![/bold yellow]")
        return

    if os.path.exists(file_name):
        workbook = load_workbook(file_name)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name

        # Add headings only for a new file
        sheet.append(headers)

    for row in data:
        sheet.append([str(value) if value is not None else "NULL" for value in row])

    workbook.save(file_name)

    print(f"[bold green]Data exported successfully to {file_name}![/bold green]")
